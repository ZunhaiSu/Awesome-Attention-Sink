"""Import the supplied Feishu export without reassessing scholarly labels.

One-time importer: pip install openpyxl. Validation and export use only stdlib.
"""
import argparse
import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
HEADERS = ['Label', 'DATE', 'Paper', 'Link', 'Pub', 'Models', 'Utilization',
           'Understanding', 'Mitigation', 'Application']
CORE = {
    'utilization': ['Sink Token Preservation', 'Attention Redistribution',
                    'Learnable Prefix Tokens', 'Sink Token Repurposing'],
    'interpretation': ['Softmax Limitations and No-Op Theory', 'Outlier Circuits',
                       'Implicit Attention Bias', 'Geometric Anchoring',
                       'Other Mechanistic Interpretations'],
    'mitigation': ['Gated Attention Mechanisms', 'Modified Softmax Functions',
                   'Learnable Attention Bias', 'Pre-training Interventions',
                   'Other Mitigation Techniques'],
}
APPS = ['Model Pre-training', 'Model Tuning', 'Efficient Inference',
        'Mechanism Interpretability', 'Reducing Hallucination', 'Safety & Robustness',
        'General Capability Enhancement', 'Long-Context Enhancement',
        'Multi-Modal Enhancement']
MODELS = {
    '2.3.1': 'Classical Language Models',
    '2.3.2': 'Large Language Models',
    '2.3.3': 'Mixture-of-Experts Large Language Models',
    '2.3.4': 'Multi-Modal Large Language Models',
    '2.4': 'Vision Transformers',
    '2.5': 'Diffusion Transformers',
}
TITLES = {'utilization': 'Fundamental Utilization',
          'interpretation': 'Mechanistic Interpretation',
          'mitigation': 'Strategic Mitigation'}


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')


def split_tags(value):
    if value is None or str(value).strip() in {'', 'N/A'}:
        return []
    return [s.strip() for s in str(value).split(',')]


def unique(values):
    return list(dict.fromkeys(values))


def month_date(value):
    text = str(value)
    if not re.fullmatch(r'\d{4}', text):
        raise ValueError(f'Unsupported legacy date: {value!r}')
    year, month = 2000 + int(text[:2]), int(text[2:])
    # The source uses YYMM; 1906 means June 2019, not the year 1906.
    if month == 0:
        return {'value': str(year), 'precision': 'year'}
    if not 1 <= month <= 12:
        raise ValueError(f'Invalid legacy month: {value!r}')
    return {'value': f'{year:04d}-{month:02d}', 'precision': 'month'}


def identify(url):
    u = urlsplit(url)
    if u.scheme not in {'http', 'https'}:
        raise ValueError(f'Unsupported paper URL: {url}')
    host = u.hostname or ''
    if host in {'arxiv.org', 'www.arxiv.org', 'export.arxiv.org'}:
        m = re.search(r'/(?:abs|pdf)/(\d{4}\.\d{4,5})(v\d+)?', u.path)
        if m:
            return 'as-arxiv-' + m[1], {'arxiv': m[1]}, 'https://arxiv.org/abs/' + m[1]
    if host == 'openreview.net':
        key = parse_qs(u.query).get('id', [None])[0]
        if key:
            return 'as-openreview-' + key, {'openreview': key}, 'https://openreview.net/forum?id=' + key
    doi = re.search(r'10\.\d{4,9}/[^?#]+', url)
    ids = {'doi': doi[0]} if doi else {}
    return 'as-link-' + hashlib.sha256(url.encode()).hexdigest()[:16], ids, url


def canonical(axis, tag):
    if axis in CORE:
        section = {'utilization': 3, 'interpretation': 4, 'mitigation': 5}[axis]
        m = re.match(rf'{section}\.(\d+)\b', tag)
        if not m or not 1 <= int(m[1]) <= len(CORE[axis]):
            raise ValueError(f'Unmapped {axis} tag: {tag}')
        return CORE[axis][int(m[1]) - 1]
    if axis == 'applications':
        m = re.match(r'(\d+)\b', tag)
        if not m or not 1 <= int(m[1]) <= len(APPS):
            raise ValueError(f'Unmapped application: {tag}')
        return APPS[int(m[1]) - 1]
    for key, label in MODELS.items():
        if tag.startswith(key + ' '):
            return label
    if tag.startswith('2.6 Others:'):
        return tag.split(':', 1)[1].strip()
    raise ValueError(f'Unmapped architecture: {tag}')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('source', type=Path)
    ap.add_argument('--date', required=True)
    args = ap.parse_args()
    wb = load_workbook(args.source, data_only=False)
    if len(wb.worksheets) != 1:
        raise ValueError('Inspect a multi-sheet export before importing.')
    sheet = wb.worksheets[0]
    assert [c.value for c in sheet[2]] == HEADERS
    assert not any(c.data_type == 'f' for row in sheet for c in row), 'Formula source needs review.'
    source_id = 'feishu-' + args.date
    raw_rows, papers, aliases = [], [], {}
    source_cols = {'architectures': 'Models', 'utilization': 'Utilization',
                   'interpretation': 'Understanding', 'mitigation': 'Mitigation',
                   'applications': 'Application'}
    for row_no in range(3, sheet.max_row + 1):
        cells = list(sheet[row_no])
        if all(c.value is None for c in cells):
            continue
        values = {h: c.value for h, c in zip(HEADERS, cells)}
        assert values['Paper'] and values['Link'], row_no
        raw_rows.append({'row': row_no, 'values': values,
                         'hyperlinks': {h: c.hyperlink.target for h, c in zip(HEADERS, cells) if c.hyperlink}})
        ident, ids, url = identify(str(values['Link']).strip())
        classification, legacy_tags = {}, {}
        for axis, col in source_cols.items():
            tags = split_tags(values[col])
            legacy_tags[axis] = tags
            classification[axis] = unique(canonical(axis, t) for t in tags)
            for tag in tags:
                aliases[(axis, tag)] = canonical(axis, tag)
        papers.append({
            'id': ident, 'title': str(values['Paper']).strip(), 'identifiers': ids,
            'date': month_date(values['DATE']),
            'publication': {'label': values['Pub'], 'verification': 'imported_not_rechecked'},
            'links': {'paper': url, 'original': str(values['Link'])},
            'classification': classification, 'legacy_tags': legacy_tags,
            'legacy_new_flag': values['Label'] == 'New',
            'verification': {'classification': 'legacy_imported', 'review_in_this_update': 'not_reassessed'},
            'provenance': {'source_id': source_id, 'sheet': sheet.title, 'row': row_no},
        })
    assert len(papers) == len({p['id'] for p in papers}), 'Resolve duplicate identifiers before import.'
    # Do not collapse uncertain title matches automatically.
    titles = [''.join(c for c in p['title'].casefold() if c.isalnum()) for p in papers]
    assert len(titles) == len(set(titles)), 'Review duplicate titles before import.'
    digest = hashlib.sha256(args.source.read_bytes()).hexdigest()
    raw = {'source_id': source_id, 'source_filename': args.source.name, 'sha256': digest,
           'captured_on': args.date, 'sheet': sheet.title, 'headers': HEADERS,
           'first_row': [c.value for c in sheet[1]], 'rows': raw_rows}
    write_json(ROOT / 'data/sources' / (source_id + '.json'), raw)
    write_json(ROOT / 'data/papers.json', papers)
    write_json(ROOT / 'data/taxonomy.json', {
        'schema_version': 1, 'core_axes': {a: {'name': TITLES[a], 'categories': cats} for a, cats in CORE.items()},
        'applications': APPS, 'architectures': sorted({t for p in papers for t in p['classification']['architectures']}),
        'aliases': [{'axis': a, 'source_label': t, 'canonical_label': v} for (a, t), v in sorted(aliases.items())],
        'rules': ['Labels may overlap. Empty imported labels mean unassigned, not evidence of absence.',
                  'Legacy fine-grained labels are retained verbatim beside canonical categories.',
                  'Taxonomy mapping does not reassess the source authors’ scholarly classification.']})
    stats = {
        'date': args.date, 'source_id': source_id, 'source_sha256': digest,
        'source_rows': len(raw_rows), 'migrated_records': len(papers),
        'arxiv_records': sum('arxiv' in p['identifiers'] for p in papers),
        'other_links': sum('arxiv' not in p['identifiers'] for p in papers),
        'legacy_new_flags': sum(p['legacy_new_flag'] for p in papers),
        'year_precision_records': sum(p['date']['precision'] == 'year' for p in papers),
        'month_precision_records': sum(p['date']['precision'] == 'month' for p in papers),
        'publication_labels_imported_not_reverified': True,
        'date_range': [min(p['date']['value'] for p in papers), max(p['date']['value'] for p in papers)],
        'axis_counts': {a: sum(bool(p['classification'][a]) for p in papers) for a in CORE},
        'category_counts': {a: dict(Counter(t for p in papers for t in p['classification'][a])) for a in source_cols},
    }
    write_json(ROOT / 'data/migration.json', stats)
    print(json.dumps({k: v for k, v in stats.items() if k != 'category_counts'}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
